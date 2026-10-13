"""Extract audit fields from the original workbook without editing/exporting it."""
import json
import sys
from openpyxl import load_workbook

def text(value):
    return None if value is None or not str(value).strip() else str(value).strip()

workbook = load_workbook(sys.argv[1], read_only=True, data_only=False)
try:
    sheet = workbook['katalog']
    rows = sheet.iter_rows()
    headers = {text(cell.value): index for index, cell in enumerate(next(rows))}
    fields = ['Numer pytania', 'Media', 'Kategorie']
    if any(field not in headers for field in fields):
        raise ValueError('Missing required source columns')
    records = []
    for row in rows:
        cells = {field: row[headers[field]] for field in fields}
        if cells['Numer pytania'].value is None:
            continue
        if any(cell.data_type == 'f' for cell in cells.values()):
            raise ValueError('Formula in source audit field; cached result cannot certify reference')
        value = cells['Numer pytania'].value
        qid = int(value)
        if isinstance(value, bool) or float(value) != qid or qid <= 0:
            raise ValueError('Invalid source question ID')
        records.append({'id': qid, 'media_filename': text(cells['Media'].value),
                        'categories': [item.strip() for item in (text(cells['Kategorie'].value) or '').split(',') if item.strip()],
                        'sourceRow': cells['Numer pytania'].row, 'mediaCell': cells['Media'].coordinate})
    print(json.dumps({'sheet': 'katalog', 'records': records}, ensure_ascii=True))
finally:
    workbook.close()
